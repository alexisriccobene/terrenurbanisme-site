# Génère data/communes/{dep}.json : prix DVF par commune (Etalab) et classement loi Montagne (ANCT)
import csv, io, json, os, re, sys, urllib.request
from collections import defaultdict
import openpyxl

UA = {'User-Agent': 'TerrenUrbanisme/1.0 (contact@terrenurbanisme.com)'}
def get(url):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=300) as r:
        return r.read()
INSEE = re.compile(r'^(\d{5}|2[AB]\d{3})$')
def dep(c): return c[:3] if c.startswith('97') else c[:2]
out = defaultdict(dict); meta = {}
def note(msg): print('::notice::' + str(msg).replace('\n', ' | ')[:900], flush=True)
import traceback, atexit
def bilan():
    note('bilan ' + json.dumps({k: (v if k != 'montagne' else {r: {x: y for x, y in i.items() if x != 'exemple'} for r, i in v.items()}) for k, v in meta.items()}, ensure_ascii=False))
    for r, i in meta.get('montagne', {}).items():
        if 'exemple' in i: note('exemple ' + r[:8] + ' ' + json.dumps(i['exemple'][:6], ensure_ascii=False))
atexit.register(bilan)

# 1. Prix : statistiques DVF par commune
try:
    raw = get('https://www.data.gouv.fr/api/1/datasets/r/851d342f-9c96-41c1-924a-11a7a7aae8a6')
except Exception as e:
    note('prix téléchargement ' + repr(e)); raise
txt = raw.decode('utf-8-sig', 'replace')
sep = ';' if txt[:2000].count(';') > txt[:2000].count(',') else ','
rd = csv.DictReader(io.StringIO(txt), delimiter=sep)
meta['prix_colonnes'] = rd.fieldnames
n = 0
def num(v):
    try: return round(float(v)) if v not in (None, '', 'NA') else None
    except: return None
for r in rd:
    if (r.get('echelle_geo') or '') != 'commune': continue
    c = (r.get('code_geo') or '').strip()
    if not INSEE.match(c): continue
    p = []
    for t in ('maison', 'appartement', 'local'): p += [num(r.get('med_prix_m2_whole_' + t)), num(r.get('nb_ventes_whole_' + t)), num(r.get('moy_prix_m2_whole_' + t))]
    if any(x for x in p): out[dep(c)].setdefault(c, {})['p'] = p; n += 1
meta['prix_communes'] = n

# 2. Loi Montagne : communes classées (COG), colonne INSEE détectée
mont = {}
for rid in ['fa1eeb72-6adb-4efd-a1a1-9f1266c78315', '87bc6d48-f1ed-4924-be55-0142660033de', '8d0ae2b1-9b8c-48df-89b9-6334b166e5c5']:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(get('https://www.data.gouv.fr/api/1/datasets/r/' + rid)), read_only=True, data_only=True)
        info = {'feuilles': wb.sheetnames}
        best = None
        for ws in wb.worksheets:
            rows = [[('' if v is None else str(v).strip()) for v in row] for row in ws.iter_rows(values_only=True)]
            if not rows: continue
            w = max(len(r) for r in rows)
            for j in range(w):
                col = [r[j] if j < len(r) else '' for r in rows]
                hits = sum(1 for v in col if INSEE.match(v) or (v.isdigit() and 1001 <= int(v) <= 99999 and len(v) == 4))
                if not best or hits > best[0]: best = (hits, ws.title, j, rows)
        hits, title, j, rows = best
        info.update({'feuille': title, 'colonne': j, 'lignes_code': hits, 'exemple': rows[:6]})
        # colonne décrivant un classement partiel éventuel
        pj = None
        for k in range(max(len(r) for r in rows)):
            vals = ' '.join((r[k] if k < len(r) else '') for r in rows[:4000]).lower()
            if 'partiel' in vals: pj = k; break
        info['colonne_partiel'] = pj
        cur = {}
        for r in rows:
            v = r[j] if j < len(r) else ''
            if v.isdigit() and len(v) == 4: v = '0' + v
            if INSEE.match(v): cur[v] = 2 if (pj is not None and pj < len(r) and 'partiel' in r[pj].lower()) else 1
        info['communes'] = len(cur)
        meta.setdefault('montagne', {})[rid] = info
        if len(cur) > len(mont): mont = cur
    except Exception as e:
        meta.setdefault('montagne', {})[rid] = {'erreur': str(e)[:300]}
for c, v in mont.items(): out[dep(c)].setdefault(c, {})['mt'] = v
meta['montagne_communes'] = len(mont)

if meta['prix_communes'] < 20000 or meta['montagne_communes'] < 3000:
    sys.exit('données incomplètes : fichiers non écrits')
os.makedirs('data/communes', exist_ok=True)
for d, obj in out.items():
    with open(f'data/communes/{d}.json', 'w') as f: json.dump(obj, f, separators=(',', ':'))
with open('data/meta.json', 'w') as f: json.dump(meta, f, ensure_ascii=False, indent=1)
print('ok', meta['prix_communes'], meta['montagne_communes'], len(out))
